"""Discovery, classification, execution and verification over the simulated platforms.

The state shape mirrors the EraseFlow frontend (`DemoState` in frontend/src/App.tsx),
so the API can return it unchanged. Everything runs on synthetic data in a throwaway
working copy of `sim_data/`; nothing touches a real account.
"""
from __future__ import annotations

import json
import re
import shutil
import tempfile
import threading
from datetime import datetime, timezone
from pathlib import Path
from random import choices
from string import ascii_uppercase, digits

from openpyxl import load_workbook
from PIL import Image, ImageDraw

from . import seed
from .models import Platform, Record
from .subject import SARAH

PLATFORM_LABEL = {
    Platform.SLACK: "Slack", Platform.GMAIL: "Gmail", Platform.SALESFORCE: "Salesforce",
    Platform.NOTION: "Notion", Platform.DRIVE: "Google Drive",
}
KIND_LABEL = {
    "message": "Message", "thread": "Email thread", "contact": "Contact", "note": "Opportunity note",
    "invoice": "Financial transaction", "page": "Page", "image": "Image", "sheet": "Spreadsheet",
}
SAFE_ACTIONS = {"DELETE", "REDACT", "RETAIN"}


def now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z"


def _event(type_: str, summary: str, at: str | None = None) -> dict:
    return {"id": f"evt-{''.join(choices(ascii_uppercase + digits, k=8))}", "type": type_,
            "at": at or now(), "summary": summary}


class Engine:
    def __init__(self, sim_root: Path = Path("sim_data")) -> None:
        self.sim_root = sim_root
        self.lock = threading.Lock()
        self.work: Path | None = None
        self.records: list[Record] = []
        self.state: dict = {}
        self.reset()

    # ---------- lifecycle
    def reset(self) -> dict:
        with self.lock:
            if not (self.sim_root / "expected_labels.json").exists():
                seed.build(self.sim_root)
            if self.work:
                shutil.rmtree(self.work, ignore_errors=True)
            self.work = Path(tempfile.mkdtemp(prefix="erasure-sim-"))
            shutil.copytree(self.sim_root, self.work, dirs_exist_ok=True)
            self.records = [Record.model_validate_json(f.read_text())
                            for plat in Platform for f in sorted((self.work / plat.value).glob("*.json"))]
            self.originals = {r.id: r.model_copy(deep=True) for r in self.records}
            self.original_text = {r.id: self.text_of(r) for r in self.records}
            req = {"id": "ER-2025-0142", "name": SARAH.full_name, "email": SARAH.emails[0],
                   "phone": SARAH.phones[0], "aliases": SARAH.aliases[:1], "roles": SARAH.roles,
                   "submittedAt": "2025-02-14T09:42:00.000Z", "status": "Discovery complete"}
            matches = self.find_matches(req)
            req["recordsFound"] = len(matches)
            self.state = {
                "requests": [req], "currentRequestId": req["id"], "matches": matches,
                "decisions": {}, "execution": [], "verification": None,
                "audit": [_event("Discovery scan", self._scan_summary(matches), req["submittedAt"]),
                          _event("Request received", f"Erasure request {req['id']} created.", req["submittedAt"])],
                "crmSyncUsed": False, "scanComplete": True, "scanRunning": False,
            }
            return self.snapshot()

    def snapshot(self) -> dict:
        by_id = {m["recordId"]: m for m in self.state["matches"]}
        originals = [{**self._view(o, None), "content": self.original_text[o.id]} for o in self.originals.values()]
        return {**self.state, "records": [self._view(r, by_id.get(r.id)) for r in self.records], "originals": originals,
                "platforms": [{"name": PLATFORM_LABEL[p], "type": t, "records": sum(r.platform == p for r in self.records)}
                              for p, t in [(Platform.SLACK, "Messaging"), (Platform.GMAIL, "Email"),
                                           (Platform.SALESFORCE, "CRM"), (Platform.NOTION, "Docs"),
                                           (Platform.DRIVE, "Files")]]}

    def _scan_summary(self, matches: list[dict]) -> str:
        return f"Scanned {len(self.records)} synthetic records across {len(PLATFORM_LABEL)} systems; {len(matches)} potential matches."

    def _request(self) -> dict:
        s = self.state
        return next(r for r in s["requests"] if r["id"] == s["currentRequestId"])

    def _set_status(self, status: str) -> None:
        self._request()["status"] = status

    # ---------- record text
    def text_of(self, r: Record) -> str:
        if r.kind == "image":
            ocr = " · ".join(r.metadata.get("ocr_text", [])) or "(no text)"
            return f"OCR text: {ocr}" + (" · face detected" if r.metadata.get("has_face") else "")
        if r.kind == "sheet":
            ws = load_workbook(self.work / r.file).active
            return "\n".join(" | ".join(str(c) for c in row) for row in ws.iter_rows(values_only=True) if any(row))
        return r.content

    def _view(self, r: Record, m: dict | None) -> dict:
        return {"id": r.id, "platform": PLATFORM_LABEL[r.platform], "type": KIND_LABEL.get(r.kind, r.kind),
                "title": r.title, "content": self.text_of(r), "createdAt": r.created_at,
                "evidence": m["evidence"] if m else "No requester identifiers found",
                "suggestedAction": m["suggestedAction"] if m else "REJECT",
                "confidence": m["confidence"] if m else 1, "ambiguity": m["ambiguity"] if m else False,
                "rationale": m["rationale"] if m else "No email, phone, full-name, alias or role evidence links this record to the requester."}

    # ---------- discovery + classification
    def classify(self, r: Record, req: dict) -> tuple[str, str]:
        if r.kind in {"invoice"}:
            return "RETAIN", "Identity is clear, but this is a financial record that needs a documented retention review."
        if (r.owner or "").lower() in {req["email"].lower(), (SARAH.crm_contact_id or "").lower()}:
            return "DELETE", "The item belongs to the requester, so the whole item is erased."
        if r.kind == "sheet":
            return "DELETE", "Only the requester's row is removed; other people's rows stay."
        return "REDACT", "The requester is mentioned inside an item that belongs to someone else; redact only the identifiers and keep the rest."

    def find_matches(self, req: dict) -> list[dict]:
        name = req["name"].lower().strip()
        email = req["email"].lower().strip()
        phone = re.sub(r"\D", "", req["phone"])
        aliases = [a.strip() for a in req["aliases"] if len(a.strip()) > 1]
        roles = [x.strip() for x in req.get("roles", []) if x.strip()]
        out = []
        for r in self.records:
            text = self.text_of(r)
            low = text.lower()
            email_hit = len(email) > 3 and email in low
            phone_hit = len(phone) >= 7 and phone in re.sub(r"\D", "", text)
            name_hit = len(name) > 2 and name in low
            alias = next((a for a in aliases if re.search(rf"\b{re.escape(a)}\b", text, re.I)), None)
            role = next((x for x in roles if x.lower() in low), None)
            if not (email_hit or phone_hit or name_hit or alias or role):
                continue
            ambiguity = bool(alias) and not (email_hit or phone_hit or name_hit or role)
            evidence = " · ".join(filter(None, [
                "Email match" if email_hit else None, "Phone match" if phone_hit else None,
                "Full name match" if name_hit else None,
                f"Alias “{alias}”{' only' if ambiguity else ''}" if alias else None,
                f"Role-only reference “{role}” (indirect)" if role else None]))
            if ambiguity:
                action, why, conf = "MANUAL REVIEW", "An alias-only match is not proof of identity. Keep this item unresolved until a reviewer can verify it.", 41
            else:
                action, why = self.classify(r, req)
                conf = 72 if (role and not (email_hit or phone_hit or name_hit)) else min(99, 78 + 17 * email_hit + 15 * phone_hit + 12 * name_hit)
                if role and not (email_hit or phone_hit or name_hit):
                    why += " Matched by role only; a reviewer should confirm the role resolves to the requester."
            out.append({"recordId": r.id, "evidence": evidence, "suggestedAction": action,
                        "confidence": conf, "ambiguity": ambiguity, "rationale": why})
        return out

    # ---------- redaction helpers
    def _redact_tokens(self, req: dict) -> list[str]:
        toks = [req["email"], req["phone"], req["name"], *req["name"].split(), *req["aliases"], *req.get("roles", [])]
        return sorted({t for t in toks if t}, key=len, reverse=True)

    def _redact_text(self, text: str, req: dict) -> str:
        for t in self._redact_tokens(req):
            text = re.sub(re.escape(t), "[REDACTED]", text, flags=re.I)
        return text

    def _mask_image(self, r: Record) -> None:
        path = self.work / r.file
        img = Image.open(path).convert("RGB")
        d = ImageDraw.Draw(img)
        d.rectangle([240, 100, 640, 200], fill="black")        # name text block
        if r.metadata.get("has_face"):
            d.rectangle([55, 85, 205, 255], fill="black")      # face region
        img.save(path)                                          # re-saving drops EXIF
        r.metadata["ocr_text"] = ["[REDACTED]"] * len(r.metadata.get("ocr_text", []))
        r.metadata["has_face"] = False

    def _delete_rows(self, r: Record, req: dict) -> int:
        wb = load_workbook(self.work / r.file)
        ws = wb.active
        keys = [req["email"].lower(), req["name"].lower()]
        rows = [i for i, row in enumerate(ws.iter_rows(values_only=True), 1)
                if i > 1 and any(k in " ".join(map(str, row)).lower() for k in keys)]
        for i in reversed(rows):
            ws.delete_rows(i)
        wb.save(self.work / r.file)
        return len(rows)

    # ---------- API operations (each returns a fresh snapshot)
    def create_request(self, data: dict) -> dict:
        with self.lock:
            rid = f"ER-{datetime.now().year}-{''.join(choices(ascii_uppercase + digits, k=5))}"
            req = {**data, "id": rid, "submittedAt": now(), "status": "New"}
            s = self.state
            s["requests"].insert(0, req)
            s["currentRequestId"] = rid
            s.update(decisions={}, execution=[], verification=None)
            s["audit"].insert(0, _event("Request received", f"Erasure request {rid} created for {req['name']}."))
            self._scan(req)
            return self.snapshot()

    def _scan(self, req: dict) -> None:
        s = self.state
        s["matches"] = self.find_matches(req)
        s.update(scanComplete=True, scanRunning=False, verification=None)
        req["status"] = "Discovery complete"
        req["recordsFound"] = len(s["matches"])
        s["audit"].insert(0, _event("Discovery scan", self._scan_summary(s["matches"])))

    def select(self, rid: str) -> dict:
        with self.lock:
            s = self.state
            if not any(r["id"] == rid for r in s["requests"]):
                raise KeyError(rid)
            s.update(currentRequestId=rid, decisions={}, execution=[], verification=None, scanComplete=True, scanRunning=False)
            s["matches"] = self.find_matches(self._request())
            return self.snapshot()

    def scan(self) -> dict:
        with self.lock:
            self._scan(self._request())
            return self.snapshot()

    def decide(self, d: dict) -> dict:
        with self.lock:
            s = self.state
            d = {**d, "updatedAt": now()}
            s["decisions"][d["recordId"]] = d
            s["verification"] = None
            self._set_status("In review")
            s["audit"].insert(0, _event("Reviewer decision", f"{d['recordId']}: {d['action']}{' approved' if d['approved'] else ''}."))
            return self.snapshot()

    def execute(self) -> dict:
        with self.lock:
            s, req = self.state, self._request()
            done = {e["recordId"] for e in s["execution"] if e["result"] == "success"}
            approved = [d for m in s["matches"] if (d := s["decisions"].get(m["recordId"]))
                        and d["approved"] and d["action"] in SAFE_ACTIONS and d["recordId"] not in done]
            results = []
            for d in approved:
                r = next((x for x in self.records if x.id == d["recordId"]), None)
                res = lambda ok, msg: results.append({"recordId": d["recordId"], "action": d["action"],
                                                      "result": "success" if ok else "failed", "at": now(), "details": msg})
                if r is None:
                    res(False, "The record was no longer available at execution time.")
                elif d["action"] == "DELETE" and r.kind == "sheet":
                    res(True, f"Removed {self._delete_rows(r, req)} row(s) for the requester; other rows preserved.")
                elif d["action"] == "DELETE":
                    self.records.remove(r)
                    res(True, "Record removed from the simulated platform.")
                elif d["action"] == "REDACT" and r.kind == "image":
                    self._mask_image(r)
                    res(True, "Name and face masked in the image; EXIF metadata stripped.")
                elif d["action"] == "REDACT":
                    r.content = self._redact_text(r.content, req)
                    res(True, "Requester identifiers redacted; remaining content preserved.")
                else:
                    res(True, f"Record retained. Rationale: {d['justification']}")
            if not results:
                return self.snapshot()
            s["execution"] = results + s["execution"]
            s["verification"] = None
            self._set_status("Executed")
            ok = sum(x["result"] == "success" for x in results)
            s["audit"].insert(0, _event("Execution", f"{ok} approved action(s) completed; {len(results) - ok} failed."))
            return self.snapshot()

    def crm_sync(self) -> dict:
        """Demo-only: a stale CRM sync re-creates one erased record, so Verify can catch it."""
        with self.lock:
            s = self.state
            if s["crmSyncUsed"]:
                return self.snapshot()
            live = {r.id for r in self.records}
            hit = next((e for e in reversed(s["execution"]) if e["action"] == "DELETE"
                        and e["result"] == "success" and e["recordId"] not in live), None)
            if hit:
                self.records.append(self.originals[hit["recordId"]].model_copy(deep=True))
                s["crmSyncUsed"] = True
                s["verification"] = None
                s["audit"].insert(0, _event("Demo sync", f"Simulated CRM sync restored deleted record {hit['recordId']}."))
            return self.snapshot()

    def verify(self) -> dict:
        with self.lock:
            s, req = self.state, self._request()
            current = {m["recordId"] for m in self.find_matches(req)}
            last = {e["recordId"]: e for e in reversed(s["execution"])}
            buckets = {k: [] for k in ("deleted", "redacted", "retained", "unresolved", "failed", "reappeared")}
            for m in s["matches"]:
                rid, d, e = m["recordId"], s["decisions"].get(m["recordId"]), last.get(m["recordId"])
                if e and e["result"] == "failed":
                    buckets["failed"].append(rid)
                elif e and e["action"] == "DELETE":
                    buckets["reappeared" if rid in current else "deleted"].append(rid)
                elif e and e["action"] == "REDACT":
                    buckets["unresolved" if rid in current else "redacted"].append(rid)
                elif e and e["action"] == "RETAIN":
                    buckets["retained"].append(rid)
                elif d and d["action"] == "REJECT" and d["justification"].strip():
                    continue
                else:
                    buckets["unresolved"].append(rid)
            incomplete = bool(buckets["unresolved"] or buckets["reappeared"] or buckets["failed"])
            v = {"status": "INCOMPLETE" if incomplete else "COMPLETE", "checkedAt": now(),
                 **{f"{k}Ids": v for k, v in buckets.items()}}
            s["verification"] = v
            self._set_status("Executed" if incomplete else "Verified")
            s["audit"].insert(0, _event("Verification", f"{v['status']}: {len(buckets['deleted'])} deletion(s), "
                                        f"{len(buckets['redacted'])} redaction(s), {len(buckets['unresolved'])} unresolved item(s)."))
            return self.snapshot()

    def certificate(self) -> dict:
        with self.lock:
            s, req = self.state, self._request()
            v = s["verification"]
            return {
                "requestId": req["id"], "subject": req["name"], "submittedAt": req["submittedAt"],
                "generatedAt": now(), "status": v["status"] if v else "NOT VERIFIED",
                "scope": [PLATFORM_LABEL[p] for p in Platform],
                "items": [{"recordId": m["recordId"], "evidence": m["evidence"], "decision": s["decisions"].get(m["recordId"]),
                           "execution": next((e for e in s["execution"] if e["recordId"] == m["recordId"]), None)}
                          for m in s["matches"]],
                "verification": v, "auditTrail": list(reversed(s["audit"])),
                "limitations": ["Simulated platforms only.",
                                "A zero-hit scan is not a compliance certificate; backups, trash and version history are not modelled.",
                                "Not legal advice."],
            }
